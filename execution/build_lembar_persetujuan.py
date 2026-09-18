r"""
build_lembar_persetujuan.py
Generates Lampiran 2: Contoh Halaman Persetujuan Proposal Tugas Akhir
according to Buku Pedoman Penyusunan Tugas Akhir UKRIDA FEB 2023 (Halaman 32).
"""

import os
import sys
import subprocess
from pathlib import Path
import docx
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import fitz

BASE_DIR = Path(r"d:\Perkuliahan\Skripsi\SKRIPSI-arthur")
OUT_DIR = BASE_DIR / "01_Naskah_Utama" / "01_Lembar_Persetujuan_Proposal"
OUT_DIR.mkdir(parents=True, exist_ok=True)

DOCX_PATH = OUT_DIR / "Halaman_Persetujuan_Proposal_Arthur_Reezan.docx"
PDF_PATH = OUT_DIR / "Halaman_Persetujuan_Proposal_Arthur_Reezan.pdf"
PNG_PATH = OUT_DIR / "Halaman_Persetujuan_Proposal_Arthur_Reezan.png"

def remove_table_borders(table):
    """Remove all visible borders from a table."""
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(tblBorders)

def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    """Set explicit margins (padding) for a table cell in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_run_font(run, name="Times New Roman", size_pt=12, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{name}" w:hAnsi="{name}" w:cs="{name}"/>')
    rPr.append(rFonts)

def create_document():
    doc = Document()

    # 1. Page Setup: ISO A4, Margins: Left 4.0 cm, Right 3.0 cm, Top 3.0 cm, Bottom 3.0 cm
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(4.0)
    section.right_margin = Cm(3.0)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(3.0)

    # 2. (Header Lampiran dihapus atas permintaan user — langsung ke judul formulir)

    # 3. Form Title: PERSETUJUAN PROPOSAL TUGAS AKHIR
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(20)
    p_title.paragraph_format.line_spacing = 1.0
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("PERSETUJUAN PROPOSAL TUGAS AKHIR")
    set_run_font(r_title, size_pt=12, bold=True)

    # 4. Identity Table (borderless)
    # Available width: 21 - 4 - 3 = 14 cm
    # Col 0: Label (3.8 cm), Col 1: Colon (0.4 cm), Col 2: Content (9.8 cm)
    id_data = [
        ("Nama", ":", [
            ("Arthur Reezan", True),
            ("                                                         Pria", False)
        ]),
        ("N.I.M.", ":", [
            ("312023002", True)
        ]),
        ("Alamat", ":", [
            ("Jl. Perumahan Duta Bandara Permai Blok gs 1 No.16, RT.4/RW.10, Jatimulya, Kec. Kosambi", False)
        ]),
        ("No. Telp", ":", [
            ("085215001339", False)
        ]),
        ("Judul yang diajukan", ":", [
            ("PENGARUH HEDONIC MOTIVATION, DESIRE FOR COMPLETENESS, DAN SPECULATIVE MOTIVE TERHADAP IMPULSIVE BUYING BOOSTER PACK KARTU POKÉMON TCG DENGAN SELF-CONTROL SEBAGAI VARIABEL MODERASI", True)
        ]),
    ]

    table_id = doc.add_table(rows=len(id_data), cols=3)
    table_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(table_id)

    col_widths = [Cm(3.8), Cm(0.4), Cm(9.8)]
    for row in table_id.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
            set_cell_margins(row.cells[idx], top=30, bottom=30, left=0, right=0)

    for r_idx, (label, sep, content_runs) in enumerate(id_data):
        row = table_id.rows[r_idx]
        
        # Col 0: Label
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.line_spacing = 1.15
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(label)
        set_run_font(r0, size_pt=12)

        # Col 1: Sep
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.line_spacing = 1.15
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(sep)
        set_run_font(r1, size_pt=12)

        # Col 2: Content
        p2 = row.cells[2].paragraphs[0]
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.15
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for text, bold in content_runs:
            r2 = p2.add_run(text)
            set_run_font(r2, size_pt=12, bold=bold)

    # 5. Date: Jakarta, 18 September 2026
    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_before = Pt(28)
    p_date.paragraph_format.space_after = Pt(14)
    p_date.paragraph_format.line_spacing = 1.15
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_date = p_date.add_run("Jakarta, 18 September 2026")
    set_run_font(r_date, size_pt=12)

    # 6. Menyetujui,
    p_menyetujui = doc.add_paragraph()
    p_menyetujui.paragraph_format.space_before = Pt(0)
    p_menyetujui.paragraph_format.space_after = Pt(14)
    p_menyetujui.paragraph_format.line_spacing = 1.15
    p_menyetujui.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_menyetujui = p_menyetujui.add_run("Menyetujui,")
    set_run_font(r_menyetujui, size_pt=12)

    # 7. Signature Table: Dosen Pembimbing & Dosen Pendamping
    # 2 columns: each 7.0 cm
    table_sig1 = doc.add_table(rows=1, cols=2)
    table_sig1.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(table_sig1)

    for cell in table_sig1.rows[0].cells:
        cell.width = Cm(7.0)
        set_cell_margins(cell, top=0, bottom=0, left=40, right=40)

    # Left cell: Dosen Pembimbing
    cell_pembimbing = table_sig1.rows[0].cells[0]
    p_pb = cell_pembimbing.paragraphs[0]
    p_pb.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pb.paragraph_format.line_spacing = 1.15
    p_pb.paragraph_format.space_before = Pt(0)
    p_pb.paragraph_format.space_after = Pt(0)
    r_pb1 = p_pb.add_run("Dosen Pembimbing\n\n\n\n\n")
    set_run_font(r_pb1, size_pt=12)
    r_pb2 = p_pb.add_run("(Dr. Fredella Colline, S.E., M.M., CFP®, PFM, CHCP-A)")
    set_run_font(r_pb2, size_pt=12, bold=True)

    # Right cell: Dosen Pendamping
    cell_pendamping = table_sig1.rows[0].cells[1]
    p_pd = cell_pendamping.paragraphs[0]
    p_pd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pd.paragraph_format.line_spacing = 1.15
    p_pd.paragraph_format.space_before = Pt(0)
    p_pd.paragraph_format.space_after = Pt(0)
    r_pd1 = p_pd.add_run("Dosen Pendamping\n\n\n\n\n")
    set_run_font(r_pd1, size_pt=12)
    r_pd2 = p_pd.add_run("( ...................................................... )")
    set_run_font(r_pd2, size_pt=12)

    # 8. Mengetahui,
    p_mengetahui = doc.add_paragraph()
    p_mengetahui.paragraph_format.space_before = Pt(18)
    p_mengetahui.paragraph_format.space_after = Pt(0)
    p_mengetahui.paragraph_format.line_spacing = 1.15
    p_mengetahui.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_menyetujui2 = p_mengetahui.add_run("Mengetahui,\n\n\n\n\n")
    set_run_font(r_menyetujui2, size_pt=12)
    r_kaprodi_sig = p_mengetahui.add_run("(Dr. Daniel Widjaja, S.E., M.M., CSCU, CEAP)\n")
    set_run_font(r_kaprodi_sig, size_pt=12, bold=True)
    r_kaprodi = p_mengetahui.add_run("Ketua Program Studi Sementara")
    set_run_font(r_kaprodi, size_pt=12)

    doc.save(str(DOCX_PATH))
    print(f"[OK] DOCX generated: {DOCX_PATH}")

def convert_docx_to_pdf():
    """Convert DOCX to PDF using PowerShell and Word COM."""
    ps_script = f"""
    $docxPath = '{str(DOCX_PATH)}'
    $pdfPath = '{str(PDF_PATH)}'
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    try {{
        $doc = $word.Documents.Open($docxPath)
        $doc.SaveAs([ref]$pdfPath, [ref]17)
        $doc.Close()
    }} finally {{
        $word.Quit()
    }}
    """
    ps_file = OUT_DIR / "_temp_convert.ps1"
    with open(ps_file, "w", encoding="utf-8") as f:
        f.write(ps_script)

    try:
        res = subprocess.run(
            ["powershell", "-ExecutionPolicy", "Bypass", "-File", str(ps_file)],
            capture_output=True,
            text=True
        )
        if res.returncode == 0 and PDF_PATH.exists():
            print(f"[OK] PDF generated: {PDF_PATH}")
        else:
            print(f"[ERROR] PowerShell failed: {res.stderr}")
    finally:
        if ps_file.exists():
            ps_file.unlink()

def generate_preview():
    """Render page 1 of PDF to PNG for visual inspection."""
    if not PDF_PATH.exists():
        print("[WARN] PDF not found for rendering preview.")
        return
    doc = fitz.open(str(PDF_PATH))
    print(f"Total PDF pages: {len(doc)}")
    assert len(doc) == 1, f"ERROR: PDF has {len(doc)} pages! Expected exactly 1 page."
    page = doc[0]
    pix = page.get_pixmap(dpi=150)
    pix.save(str(PNG_PATH))
    print(f"[OK] PNG preview saved: {PNG_PATH}")

if __name__ == "__main__":
    create_document()
    convert_docx_to_pdf()
    generate_preview()
