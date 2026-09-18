"""
build_form_bimbingan.py
Updates 00 - Form Bimbingan Skripsi.docx with Arthur Reezan's identity
and converts it to PDF using Word COM.
"""

import os
import sys
import subprocess
from pathlib import Path
import docx
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import fitz

BASE_DIR = Path(r"d:\Perkuliahan\Skripsi\SKRIPSI-arthur\01_Naskah_Utama\01_Lembar_Persetujuan_Proposal")
DOCX_IN = BASE_DIR / "00 - Form Bimbingan Skripsi.docx"
DOCX_OUT = BASE_DIR / "00 - Form Bimbingan Skripsi.docx"
PDF_OUT = BASE_DIR / "00 - Form Bimbingan Skripsi.pdf"

def set_run_font(run, name="Verdana", size_pt=10.0, bold=False, italic=False, underline=False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{name}" w:hAnsi="{name}" w:cs="{name}"/>')
    rPr.append(rFonts)

def update_identity(doc):
    # Paragraph pairs for page 1 and page 2:
    # Page 1: 5 (Nama), 6 (NIM), 7 (Dosen 1), 8 (Dosen 2)
    # Page 2: 17 (Nama), 18 (NIM), 19 (Dosen 1), 20 (Dosen 2)
    
    sections = [(5, 6, 7, 8), (17, 18, 19, 20)]
    
    nama_text = "Arthur Reezan"
    nim_text = "312023002"
    dosen1_text = "Dr. Fredella Colline, S.E., M.M., CFP®, PFM, CHCP-A"
    dosen2_text = "-"
    
    for p_nama_idx, p_nim_idx, p_d1_idx, p_d2_idx in sections:
        # 1. Nama Mahasiswa
        p_nama = doc.paragraphs[p_nama_idx]
        p_nama.clear()
        r0 = p_nama.add_run("Nama Mahasiswa")
        set_run_font(r0, name="Verdana", size_pt=10.0)
        r1 = p_nama.add_run("\t: ")
        set_run_font(r1, name="Verdana", size_pt=10.0)
        r2 = p_nama.add_run(nama_text)
        set_run_font(r2, name="Verdana", size_pt=10.0, bold=True)
        
        # 2. N I M
        p_nim = doc.paragraphs[p_nim_idx]
        p_nim.clear()
        r0 = p_nim.add_run("N I M")
        set_run_font(r0, name="Verdana", size_pt=10.0)
        r1 = p_nim.add_run("\t\t\t: ")
        set_run_font(r1, name="Verdana", size_pt=10.0)
        r2 = p_nim.add_run(nim_text)
        set_run_font(r2, name="Verdana", size_pt=10.0, bold=True)
        
        # 3. Dosen Pembimbing (1)
        p_d1 = doc.paragraphs[p_d1_idx]
        p_d1.clear()
        r0 = p_d1.add_run("Dosen Pembimbing")
        set_run_font(r0, name="Verdana", size_pt=10.0)
        r1 = p_d1.add_run("\t: (1). ")
        set_run_font(r1, name="Verdana", size_pt=10.0)
        r2 = p_d1.add_run(dosen1_text)
        set_run_font(r2, name="Verdana", size_pt=10.0, bold=True)
        
        # 4. Dosen Pembimbing (2)
        p_d2 = doc.paragraphs[p_d2_idx]
        p_d2.clear()
        r0 = p_d2.add_run("\t\t\t  (2). ")
        set_run_font(r0, name="Verdana", size_pt=10.0)
        r1 = p_d2.add_run(dosen2_text)
        set_run_font(r1, name="Verdana", size_pt=10.0)

def convert_docx_to_pdf(docx_path, pdf_path):
    ps_script = f"""
$docxPath = '{str(docx_path)}'
$pdfPath = '{str(pdf_path)}'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {{
    $doc = $word.Documents.Open($docxPath)
    $doc.SaveAs([ref]$pdfPath, [ref]17)
    $doc.Close()
}} finally {{
    $word.Quit()
}}
"""
    ps_file = BASE_DIR / "_temp_convert.ps1"
    with open(ps_file, "w", encoding="utf-8") as f:
        f.write(ps_script)

    try:
        res = subprocess.run(
            ["powershell", "-ExecutionPolicy", "Bypass", "-File", str(ps_file)],
            capture_output=True,
            text=True
        )
        if res.returncode == 0 and pdf_path.exists():
            print(f"[OK] PDF generated successfully: {pdf_path}")
        else:
            print(f"[ERROR] PowerShell failed: {res.stderr}")
    finally:
        if ps_file.exists():
            ps_file.unlink()

def render_previews(pdf_path):
    doc = fitz.open(str(pdf_path))
    print(f"Total pages in PDF: {len(doc)}")
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=150)
        png_path = BASE_DIR / f"page_{i+1}.png"
        pix.save(str(png_path))
        print(f"[OK] Rendered preview: {png_path}")

if __name__ == "__main__":
    print("Loading document...")
    doc = docx.Document(str(DOCX_IN))
    print("Updating identity...")
    update_identity(doc)
    doc.save(str(DOCX_OUT))
    print(f"[OK] Saved modified DOCX: {DOCX_OUT}")
    
    print("Converting to PDF via Word COM...")
    convert_docx_to_pdf(DOCX_OUT, PDF_OUT)
    
    print("Rendering page previews...")
    render_previews(PDF_OUT)
    print("Done!")
