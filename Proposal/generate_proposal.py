"""
generate_proposal.py
Script komprehensif untuk menyusun Proposal Skripsi (BAB 1 - BAB 3)
dalam format Markdown (.md) dan Microsoft Word (.docx)
sesuai Buku Pedoman Penyusunan Skripsi FEB Ukrida 2022 (Bab 1 Sub-bab 1.4 & Lampiran 1 & 2).

Jalankan:
z:\Skripsi\.venv\Scripts\python.exe z:\Skripsi\Proposal\generate_proposal.py
"""

import os
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─────────────────────────────────────────────────────────
# KONFIGURASI FILE & METADATA
# ─────────────────────────────────────────────────────────
BASE_DIR    = Path(r"z:\Skripsi")
SOURCE_MD   = BASE_DIR / "01_Naskah_Utama" / "SKRIPSI_ARTHUR_LENGKAP_PRISM.md"
BIB_FILE    = BASE_DIR / "latex" / "references.bib"
OUT_DIR     = BASE_DIR / "Proposal"
OUT_MD      = OUT_DIR / "PROPOSAL_SKRIPSI_ARTHUR.md"
OUT_DOCX    = OUT_DIR / "Proposal_Skripsi_Arthur.docx"

OUT_DIR.mkdir(parents=True, exist_ok=True)

JUDUL = (
    "PENGARUH PORTOFOLIO KREDIT HIJAU (GREEN FINANCING), "
    "NON-PERFORMING LOAN (NPL), DAN CAPITAL ADEQUACY RATIO (CAR) "
    "TERHADAP PROFITABILITAS (RETURN ON ASSETS / ROA) "
    "PADA BANK KBMI 4 DI INDONESIA PERIODE 2021–2025"
)
PENULIS       = "Arthur Reezan"
NIM           = "312023002"
JENIS_KELAMIN = "Pria"
PRODI         = "Program Studi Strata 1 Manajemen"
KONSENTRASI   = "Manajemen Keuangan"
PEMBIMBING    = "Dr. Diana Frederica, S.E., M.Ak., CFP®., CHCP-A"
KAPRODI       = "Rita Amelinda, S.E., M.M."
FAKULTAS      = "Fakultas Ekonomi dan Bisnis"
UNIVERSITAS   = "Universitas Kristen Krida Wacana"
KOTA          = "Jakarta"
TAHUN         = "2026"

# ─────────────────────────────────────────────────────────
# PARSER BIBTEX & FORMATTER APA
# ─────────────────────────────────────────────────────────
def parse_bibtex(bib_path: Path):
    text = bib_path.read_text(encoding="utf-8")
    entries = {}
    raw_entries = re.split(r'\n@', text)
    for raw in raw_entries:
        raw = raw.strip()
        if not raw or raw.startswith('%'):
            continue
        m = re.match(r'(\w+)\s*\{\s*([^,]+),', raw)
        if not m:
            continue
        etype = m.group(1).lower()
        key = m.group(2).strip()
        fields = {}
        for line in raw.split('\n')[1:]:
            line = line.strip()
            if '=' in line:
                k, v = line.split('=', 1)
                k = k.strip().lower()
                v = v.strip().rstrip(',').strip()
                if (v.startswith('{') and v.endswith('}')) or (v.startswith('"') and v.endswith('"')):
                    v = v[1:-1].strip()
                v = v.replace('{', '').replace('}', '')
                fields[k] = v
        entries[key] = (etype, fields)
    return entries

BIB_ENTRIES = parse_bibtex(BIB_FILE)

def format_apa_reference(key, entry):
    etype, f = entry
    author = f.get('author', '')
    year = f.get('year', '')
    title = f.get('title', '')
    
    if 'Otoritas Jasa Keuangan' in author:
        auth_str = "Otoritas Jasa Keuangan"
    elif 'Bank Indonesia' in author:
        auth_str = "Bank Indonesia"
    elif 'Universitas Kristen Krida Wacana' in author:
        auth_str = "Universitas Kristen Krida Wacana"
    else:
        authors = [a.strip() for a in author.split(' and ')]
        formatted = []
        for a in authors:
            if ',' in a:
                parts = a.split(',')
                last = parts[0].strip()
                initials = ' '.join([p.strip()[0] + '.' for p in parts[1].split() if p.strip()])
                formatted.append(f"{last}, {initials}")
            else:
                parts = a.split()
                last = parts[-1]
                initials = ' '.join([p[0] + '.' for p in parts[:-1] if p])
                formatted.append(f"{last}, {initials}")
        if len(formatted) == 1:
            auth_str = formatted[0]
        elif len(formatted) == 2:
            auth_str = f"{formatted[0]}, & {formatted[1]}"
        else:
            auth_str = ', '.join(formatted[:-1]) + f", & {formatted[-1]}"

    title = title.rstrip('.')
    if etype == 'article':
        journal = f.get('journal', '')
        volume = f.get('volume', '')
        number = f.get('number', '')
        pages = f.get('pages', '').replace('--', '–')
        doi = f.get('doi', '')
        vol_str = f"*{journal}*"
        if volume:
            vol_str += f", *{volume}*"
            if number:
                vol_str += f"({number})"
        if pages:
            vol_str += f", {pages}"
        if doi:
            vol_str += f". https://doi.org/{doi}"
        else:
            vol_str += "."
        return f"{auth_str} ({year}). {title}. {vol_str}"
    elif etype == 'book':
        edition = f.get('edition', '')
        publisher = f.get('publisher', '')
        address = f.get('address', '')
        ed_str = f" ({edition}th ed.)" if edition else ""
        pub_str = f"{address}: {publisher}." if address and publisher else f"{publisher}." if publisher else ""
        return f"{auth_str} ({year}). *{title}*{ed_str}. {pub_str}".strip()
    else:
        publisher = f.get('publisher', f.get('institution', ''))
        return f"{auth_str} ({year}). *{title}*. {publisher}."

CITE_MAP = {
    "Akomea2022": "Akomea et al. (2022)",
    "Baltagi2021": "Baltagi (2021)",
    "Barney1991": "Barney (1991)",
    "BrighamEhrhardt2020": "Brigham & Ehrhardt (2020)",
    "BrighamHouston2021": "Brigham & Houston (2021)",
    "Buallay2019": "Buallay (2019)",
    "Chiaramonte2020": "Chiaramonte et al. (2020)",
    "Cui2018": "Cui et al. (2018)",
    "Deegan2002": "Deegan (2002)",
    "Dendawijaya2015": "Dendawijaya (2015)",
    "Fachrudin2021": "Fachrudin (2021)",
    "FitrianiWardani2024": "Fitriani & Wardani (2024)",
    "Freeman2010": "Freeman (2010)",
    "GozaliSuhardi2022": "Gozali & Suhardi (2022)",
    "GujaratiPorter2020": "Gujarati & Porter (2020)",
    "GurleyShaw1960": "Gurley & Shaw (1960)",
    "HandayaniSubowo2022": "Handayani & Subowo (2022)",
    "HaryantoSudarno2023": "Haryanto & Sudarno (2023)",
    "KuncoroSuhardjono2018": "Kuncoro & Suhardjono (2018)",
    "KurniawatiYuliana2023": "Kurniawati & Yuliana (2023)",
    "LestariPurnomo2023": "Lestari & Purnomo (2023)",
    "NugrohoUtami2021": "Nugroho & Utami (2021)",
    "OJK2017": "OJK (2017)",
    "OJK2017b": "OJK (2017b)",
    "OJK2021": "OJK (2021)",
    "OJK2022": "OJK (2022)",
    "PraditaSyaichu2022": "Pradita & Syaichu (2022)",
    "Pramono2022": "Pramono et al. (2022)",
    "Riyadi2020": "Riyadi (2020)",
    "Santoso2023": "Santoso et al. (2023)",
    "SariAstuti2023": "Sari & Astuti (2023)",
    "SetiawanIndrawati2024": "Setiawan & Indrawati (2024)",
    "Spence1973": "Spence (1973)",
    "Suchman1995": "Suchman (1995)",
    "Sun2021": "Sun et al. (2021)",
    "Ukrida2022": "Universitas Kristen Krida Wacana (2022)",
    "WijayaMahardika2024": "Wijaya & Mahardika (2024)",
    "Wooldridge2020": "Wooldridge (2020)",
    "WulandariRahardjo2022": "Wulandari & Rahardjo (2022)",
    "Yin2021": "Yin et al. (2021)",
    "Zhang2022": "Zhang et al. (2022)",
    "Zhou2021": "Zhou et al. (2021)",
    "DriscollKraay1998": "Driscoll & Kraay (1998)",
    "Hoechle2007": "Hoechle (2007)",
}

def clean_citations(text: str) -> str:
    for k, v in CITE_MAP.items():
        text = re.sub(rf'\b{k}\b', v, text)
    return text

# ─────────────────────────────────────────────────────────
# LATEX TEXT CLEANER
# ─────────────────────────────────────────────────────────
def clean_latex(text: str) -> str:
    text = re.sub(r'%.*', '', text)
    text = re.sub(r'\\dfrac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', text)
    text = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', text)
    text = re.sub(r'\\left\(|\\right\)', '', text)
    text = re.sub(r'\\times', '×', text)
    text = re.sub(r'\\varepsilon', 'ε', text)
    text = re.sub(r'\\alpha', 'α', text)
    text = re.sub(r'\\beta_(\d+)', r'β\1', text)
    text = re.sub(r'\\beta', 'β', text)
    text = re.sub(r'\$([^$\n]+?)\$', lambda m: m.group(1).replace('\\', '').replace('{','').replace('}',''), text)
    text = re.sub(r'\\textbf\{([^}]*)\}', r'**\1**', text)
    text = re.sub(r'\{\\bfseries\s+([^}]*)\}', r'**\1**', text)
    text = re.sub(r'\\textit\{([^}]*)\}', r'*\1*', text)
    text = re.sub(r'\\emph\{([^}]*)\}', r'*\1*', text)
    text = re.sub(r'\\url\{([^}]*)\}', r'\1', text)
    text = re.sub(r'\\footnote\{([^}]*)\}', r' (\1)', text)
    text = re.sub(r'\\[a-zA-Z]+\*?\{[^}]*\}', '', text)
    text = re.sub(r'\\[a-zA-Z]+\*?(?:\[[^\]]*\])?', '', text)
    text = re.sub(r'[{}]', '', text)
    text = re.sub(r'  +', ' ', text)
    text = text.replace('\\\\', ' ')
    text = clean_citations(text)
    return text.strip()


def is_latex_noise(line: str) -> bool:
    line = line.strip()
    patterns = [
        r'^\\(pagenumbering|pagestyle|setcounter|addcontentsline|addtocontents)',
        r'^\\(renewcommand|tableofcontents|listoftables|listoffigures)',
        r'^\\(thispagestyle|vspace|vfill|hspace|hfill|newpage|clearpage)',
        r'^\\(begin|end)\{(titlepage|center|flushright|flushleft|document)',
        r'^\\label\{',
        r'^\\addtocounter',
        r'^---+$',
        r'^% [=\-]{5,}',
        r'^% ––',
        r'^% (====|BAGIAN|BAB)',
    ]
    return any(re.match(p, line) for p in patterns)

# ─────────────────────────────────────────────────────────
# WORD FORMATTING HELPERS
# ─────────────────────────────────────────────────────────
def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), name)
    rPr.insert(0, rFonts)


def para_fmt(para, space_before=0, space_after=0, line_spacing=1.5,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             first_indent=None, left_indent=None):
    para.alignment = align
    pf = para.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after  = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    if first_indent is not None:
        pf.first_line_indent = first_indent
    if left_indent is not None:
        pf.left_indent = left_indent


def add_runs_with_inline(para, text: str, base_size=12, base_bold=False):
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*)')
    parts = pattern.split(text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = para.add_run(part[2:-2])
            set_run_font(run, size=base_size, bold=True)
        elif part.startswith('*') and part.endswith('*'):
            run = para.add_run(part[1:-1])
            set_run_font(run, size=base_size, italic=True, bold=base_bold)
        else:
            if part:
                run = para.add_run(part)
                set_run_font(run, size=base_size, bold=base_bold)


def add_body(doc, text, first_indent=True):
    para = doc.add_paragraph()
    para_fmt(para, first_indent=Cm(1.25) if first_indent else None)
    add_runs_with_inline(para, text)
    return para


def add_heading_bab(doc, number, title):
    para = doc.add_paragraph()
    para_fmt(para, space_before=0, space_after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    line = f"BAB {number}\n{title.upper()}" if number else title.upper()
    run = para.add_run(line)
    set_run_font(run, size=12, bold=True)
    return para


def add_sub(doc, text, level=2):
    para = doc.add_paragraph()
    sp_before = 12 if level == 2 else 9
    para_fmt(para, space_before=sp_before, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    run = para.add_run(text)
    set_run_font(run, size=12, bold=True)
    return para


def add_numbered_item(doc, number, text):
    para = doc.add_paragraph()
    para_fmt(para, space_before=0, space_after=3,
             first_indent=Cm(-1.25), left_indent=Cm(1.25))
    run_num = para.add_run(f"{number}. ")
    set_run_font(run_num, size=12)
    add_runs_with_inline(para, text)
    return para


def add_bullet_item(doc, text):
    para = doc.add_paragraph()
    para_fmt(para, space_before=0, space_after=3,
             first_indent=Cm(-0.75), left_indent=Cm(1.75))
    run_bul = para.add_run("• ")
    set_run_font(run_bul, size=12)
    add_runs_with_inline(para, text)
    return para


def add_blank(doc):
    para = doc.add_paragraph()
    para_fmt(para, space_before=0, space_after=0)
    para.add_run("")
    return para


def add_center(doc, text, bold=False, size=12, space_before=0, space_after=6):
    para = doc.add_paragraph()
    para_fmt(para, space_before=space_before, space_after=space_after,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_runs_with_inline(para, text, base_size=size, base_bold=bold)
    return para


def add_right(doc, text, bold=False, size=12, space_before=0, space_after=6):
    para = doc.add_paragraph()
    para_fmt(para, space_before=space_before, space_after=space_after,
             align=WD_ALIGN_PARAGRAPH.RIGHT)
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return para


def set_margins(section):
    section.left_margin   = Cm(4)
    section.right_margin  = Cm(3)
    section.top_margin    = Cm(3)
    section.bottom_margin = Cm(3)


def add_footer_roman(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.clear()
    run = para.add_run()
    set_run_font(run, size=10)
    for tag, txt in [('begin', r' PAGE \* lowerroman '), ('end', '')]:
        fld = OxmlElement('w:fldChar')
        fld.set(qn('w:fldCharType'), tag)
        if tag == 'begin':
            instr = OxmlElement('w:instrText')
            instr.text = txt
            run._r.append(fld)
            run._r.append(instr)
        else:
            run._r.append(fld)


def add_footer_arabic(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para.clear()
    run_pg = para.add_run()
    for tag in ['begin', 'end']:
        fld = OxmlElement('w:fldChar')
        fld.set(qn('w:fldCharType'), tag)
        if tag == 'begin':
            instr = OxmlElement('w:instrText')
            instr.text = ' PAGE '
            run_pg._r.append(fld)
            run_pg._r.append(instr)
        else:
            run_pg._r.append(fld)
    set_run_font(run_pg, size=10, bold=True)
    run_uni = para.add_run("  Universitas Kristen Krida Wacana")
    set_run_font(run_uni, size=10, bold=True)


def set_table_styling(table):
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), 'CCCCCC')
        tblBorders.append(b)
    tblPr.append(tblBorders)


def set_cell_shading(cell, color_hex="F2F5F8"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

# ─────────────────────────────────────────────────────────
# FRONT MATTER BUILDERS
# ─────────────────────────────────────────────────────────
def build_judul_proposal(doc):
    """Lampiran 1: Contoh Halaman Judul Proposal Skripsi"""
    add_center(doc, "PROPOSAL SKRIPSI", bold=True, size=14, space_before=24, space_after=36)
    add_center(doc, JUDUL, bold=True, size=13, space_before=0, space_after=36)
    add_center(doc, "Diajukan kepada Program Studi Manajemen\nUntuk Menyusun Skripsi S1",
               size=12, space_before=0, space_after=36)
    add_center(doc, "Diajukan Oleh:", size=12, space_before=0, space_after=6)
    add_center(doc, PENULIS, bold=True, size=13, space_before=0, space_after=3)
    add_center(doc, f"({NIM})", bold=True, size=13, space_before=0, space_after=48)
    add_center(doc, f"FAKULTAS EKONOMI DAN BISNIS\n{UNIVERSITAS.upper()}\n{KOTA.upper()} {TAHUN}",
               bold=True, size=12, space_before=0, space_after=0)


def build_persetujuan_proposal(doc):
    """Lampiran 2: Contoh Halaman Persetujuan Proposal Skripsi"""
    add_heading_bab(doc, "", "PERSETUJUAN PROPOSAL SKRIPSI")
    add_blank(doc)
    
    info_table = doc.add_table(rows=5, cols=2)
    info_data = [
        ("Nama", f": {PENULIS}   ({JENIS_KELAMIN})"),
        ("N.I.M.", f": {NIM}"),
        ("Program Studi", f": {PRODI}"),
        ("Konsentrasi", f": {KONSENTRASI}"),
        ("Judul yang Diajukan", f": {JUDUL}"),
    ]
    for i, (k, v) in enumerate(info_data):
        c0 = info_table.rows[i].cells[0]
        c1 = info_table.rows[i].cells[1]
        c0.text = k
        c1.text = v
        for c in (c0, c1):
            for p in c.paragraphs:
                para_fmt(p, space_before=2, space_after=2)
                for r in p.runs:
                    set_run_font(r, size=12)
        c0.paragraphs[0].runs[0].bold = True

    add_blank(doc)
    add_body(doc,
        "Proposal skripsi ini telah diperiksa dan disetujui untuk dipertahankan di hadapan "
        f"Tim Penguji Seminar Proposal Skripsi {FAKULTAS} {UNIVERSITAS}.",
        first_indent=False)
    
    add_blank(doc)
    add_right(doc, f"{KOTA}, ________________ {TAHUN}", size=12, space_before=12, space_after=6)
    add_right(doc, "Menyetujui,", size=12, space_before=0, space_after=12)
    
    sig_table = doc.add_table(rows=1, cols=2)
    c_pemb = sig_table.rows[0].cells[0]
    c_pend = sig_table.rows[0].cells[1]
    
    p_pemb = c_pemb.paragraphs[0]
    p_pemb.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_pemb = p_pemb.add_run(f"Dosen Pembimbing\n\n\n\n\n( {PEMBIMBING} )")
    set_run_font(r_pemb, size=12)
    
    p_pend = c_pend.paragraphs[0]
    p_pend.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_pend = p_pend.add_run("Dosen Pendamping\n\n\n\n\n( ________________________ )")
    set_run_font(r_pend, size=12)
    
    add_blank(doc)
    add_center(doc, f"Mengetahui,\nKetua Program Studi Manajemen\n\n\n\n( {KAPRODI} )", size=12)


def build_daftar_isi_proposal(doc):
    add_heading_bab(doc, "", "DAFTAR ISI")
    add_blank(doc)
    
    entries = [
        ("HALAMAN JUDUL PROPOSAL", "i", True, 0),
        ("HALAMAN PERSETUJUAN PROPOSAL", "ii", True, 0),
        ("DAFTAR ISI", "iii", True, 0),
        ("DAFTAR TABEL", "iv", True, 0),
        ("DAFTAR GAMBAR", "v", True, 0),
        ("", "", False, 0),
        ("BAB 1  PENDAHULUAN", "1", True, 0),
        ("1.1. Latar Belakang Penelitian", "1", False, 1),
        ("1.2. Perumusan Masalah", "7", False, 1),
        ("1.3. Tujuan Penelitian", "8", False, 1),
        ("1.4. Manfaat Penelitian", "9", False, 1),
        ("1.5. Sistematika Penulisan Proposal Skripsi", "10", False, 1),
        ("", "", False, 0),
        ("BAB 2  TINJAUAN PUSTAKA DAN PENGEMBANGAN HIPOTESIS", "12", True, 0),
        ("2.1. Teori-Teori", "12", False, 1),
        ("  2.1.1. Teori Pemangku Kepentingan (Stakeholder Theory)", "12", False, 2),
        ("  2.1.2. Teori Legitimasi (Legitimacy Theory)", "14", False, 2),
        ("  2.1.3. Teori Intermediasi Finansial", "15", False, 2),
        ("  2.1.4. Teori Sinyal (Signaling Theory)", "16", False, 2),
        ("  2.1.5. Konsep Profitabilitas Bank dan Return on Assets (ROA)", "17", False, 2),
        ("  2.1.6. Portofolio Kredit Hijau (Green Financing)", "18", False, 2),
        ("  2.1.7. Risiko Kredit dan Non-Performing Loan (NPL)", "20", False, 2),
        ("  2.1.8. Kecukupan Modal dan Capital Adequacy Ratio (CAR)", "22", False, 2),
        ("2.2. Penelitian Sebelumnya", "23", False, 1),
        ("2.3. Rerangka Penelitian", "30", False, 1),
        ("2.4. Pengembangan Hipotesis", "32", False, 1),
        ("  2.4.1. Pengaruh Green Financing terhadap Return on Assets", "32", False, 2),
        ("  2.4.2. Pengaruh Non-Performing Loan terhadap Return on Assets", "34", False, 2),
        ("  2.4.3. Pengaruh Capital Adequacy Ratio terhadap Return on Assets", "36", False, 2),
        ("  2.4.4. Pengaruh Simultan Green Financing, NPL, dan CAR terhadap ROA", "37", False, 2),
        ("", "", False, 0),
        ("BAB 3  METODE PENELITIAN", "39", True, 0),
        ("3.1. Jenis, Sumber, dan Teknik Pengumpulan Data", "39", False, 1),
        ("3.2. Populasi dan Sampel", "41", False, 1),
        ("3.3. Model Penelitian", "44", False, 1),
        ("3.4. Operasionalisasi Variabel", "45", False, 1),
        ("3.5. Metode Analisis Data", "47", False, 1),
        ("  3.5.1. Analisis Statistik Deskriptif", "47", False, 2),
        ("  3.5.2. Uji Stasioneritas Panel", "48", False, 2),
        ("  3.5.3. Penentuan Model Estimasi Regresi Panel", "49", False, 2),
        ("  3.5.4. Uji Asumsi Klasik", "51", False, 2),
        ("  3.5.5. Prosedur Estimasi EViews dan Robust Standard Errors", "53", False, 2),
        ("  3.5.6. Pengujian Hipotesis Statistik", "54", False, 2),
        ("", "", False, 0),
        ("DAFTAR PUSTAKA", "57", True, 0),
    ]

    for title, page, is_bold, indent_level in entries:
        if not title:
            add_blank(doc)
            continue
        para = doc.add_paragraph()
        ind = Cm(indent_level * 0.75) if indent_level else None
        para_fmt(para, space_before=1, space_after=1, first_indent=ind)
        run_t = para.add_run(title)
        set_run_font(run_t, size=12, bold=is_bold)
        
        total_len = len(title) + len(page) + (indent_level * 3)
        dots_count = max(4, 75 - total_len)
        run_dots = para.add_run(" " + ". " * (dots_count // 2) + " ")
        set_run_font(run_dots, size=10, color=(160, 160, 160))
        
        run_p = para.add_run(page)
        set_run_font(run_p, size=12, bold=is_bold)


def build_daftar_tabel_proposal(doc):
    add_heading_bab(doc, "", "DAFTAR TABEL")
    add_blank(doc)
    
    tables = [
        ("Tabel 2.1", "Matriks Sintesis Penelitian Terdahulu (Bagian 1: Penelitian 1–5)", "24"),
        ("Tabel 2.2", "Matriks Sintesis Penelitian Terdahulu (Bagian 2: Penelitian 6–10)", "26"),
        ("Tabel 2.3", "Matriks Sintesis Penelitian Terdahulu (Bagian 3: Penelitian 11–15)", "28"),
        ("Tabel 3.1", "Daftar Entitas Sampel Penelitian Bank KBMI 4", "43"),
        ("Tabel 3.2", "Matriks Operasionalisasi Variabel Penelitian", "45"),
    ]
    for num, title, page in tables:
        para = doc.add_paragraph()
        para_fmt(para, space_before=3, space_after=3)
        run_n = para.add_run(f"{num}  ")
        set_run_font(run_n, size=12, bold=True)
        run_t = para.add_run(title)
        set_run_font(run_t, size=12)
        
        dots_len = max(4, 75 - len(num) - len(title) - len(page) - 4)
        run_dots = para.add_run(" " + ". " * (dots_len // 2) + " ")
        set_run_font(run_dots, size=10, color=(160, 160, 160))
        
        run_p = para.add_run(page)
        set_run_font(run_p, size=12, bold=True)


def build_daftar_gambar_proposal(doc):
    add_heading_bab(doc, "", "DAFTAR GAMBAR")
    add_blank(doc)
    
    figures = [
        ("Gambar 2.1", "Rerangka Pemikiran Hubungan Kausalitas Variabel Penelitian", "31"),
    ]
    for num, title, page in figures:
        para = doc.add_paragraph()
        para_fmt(para, space_before=3, space_after=3)
        run_n = para.add_run(f"{num}  ")
        set_run_font(run_n, size=12, bold=True)
        run_t = para.add_run(title)
        set_run_font(run_t, size=12)
        
        dots_len = max(4, 75 - len(num) - len(title) - len(page) - 4)
        run_dots = para.add_run(" " + ". " * (dots_len // 2) + " ")
        set_run_font(run_dots, size=10, color=(160, 160, 160))
        
        run_p = para.add_run(page)
        set_run_font(run_p, size=12, bold=True)


# ─────────────────────────────────────────────────────────
# TABLE PARSER & WRITER
# ─────────────────────────────────────────────────────────
def add_word_table_from_rows(doc, caption_text, headers, rows, source_text=""):
    para_cap = doc.add_paragraph()
    para_fmt(para_cap, space_before=12, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    run_cap = para_cap.add_run(caption_text)
    set_run_font(run_cap, size=11, bold=True)

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    set_table_styling(table)

    # Header Row
    for col_idx, h_text in enumerate(headers):
        cell = table.rows[0].cells[col_idx]
        set_cell_shading(cell, "EBF3FB")
        p = cell.paragraphs[0]
        para_fmt(p, space_before=4, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(h_text)
        set_run_font(r, size=10, bold=True)

    # Data Rows
    for row_idx, r_data in enumerate(rows):
        for col_idx, val in enumerate(r_data):
            if col_idx < len(headers):
                cell = table.rows[row_idx + 1].cells[col_idx]
                p = cell.paragraphs[0]
                align = WD_ALIGN_PARAGRAPH.CENTER if (col_idx == 0 or len(val) <= 10) else WD_ALIGN_PARAGRAPH.JUSTIFY
                para_fmt(p, space_before=3, space_after=3, line_spacing=1.15, align=align)
                add_runs_with_inline(p, val, base_size=10)

    if source_text:
        para_src = doc.add_paragraph()
        para_fmt(para_src, space_before=4, space_after=12, align=WD_ALIGN_PARAGRAPH.LEFT)
        run_src = para_src.add_run(source_text)
        set_run_font(run_src, size=10, italic=True)


def convert_latex_tables_to_markdown(text: str) -> str:
    """Converts LaTeX tabularx tables into clean GitHub Flavored Markdown tables."""
    def table_replacer(match):
        table_raw = match.group(0)
        cap_m = re.search(r'\\caption\{([^}]+)\}', table_raw)
        caption = cap_m.group(1) if cap_m else "Tabel"
        
        src_m = re.search(r'\\begin\{flushleft\}\s*\*?([^*]+)\*?\s*\\end\{flushleft\}', table_raw)
        source = src_m.group(1).strip() if src_m else ""
        
        tab_m = re.search(r'\\begin\{tabularx\}\s*\{[^}]*\}\s*\{[^}]*\}(.*?)\\end\{tabularx\}', table_raw, re.DOTALL)
        if not tab_m:
            return ""
        
        content = tab_m.group(1)
        row_lines = [rl.strip() for rl in content.split(r'\\') if '&' in rl]
        parsed_rows = []
        for rl in row_lines:
            rl = re.sub(r'\\(hline|toprule|midrule|bottomrule)', '', rl).strip()
            if not rl:
                continue
            cells = [clean_latex(c).replace('|', '/').replace('\n', ' ') for c in rl.split('&')]
            parsed_rows.append(cells)
            
        if not parsed_rows:
            return ""
            
        headers = parsed_rows[0]
        data_rows = parsed_rows[1:]
        
        md_table_str = f"\n\n**{caption}**\n\n"
        md_table_str += "| " + " | ".join(headers) + " |\n"
        md_table_str += "| " + " | ".join(["---"] * len(headers)) + " |\n"
        for dr in data_rows:
            while len(dr) < len(headers):
                dr.append("")
            md_table_str += "| " + " | ".join(dr[:len(headers)]) + " |\n"
            
        if source:
            md_table_str += f"\n*{source}*\n\n"
            
        return md_table_str

    return re.sub(r'\\begin\{table\}.*?\\end\{table\}', table_replacer, text, flags=re.DOTALL)


# ─────────────────────────────────────────────────────────
# MAIN BUILDER FOR DOCX & MARKDOWN
# ─────────────────────────────────────────────────────────
def run_proposal_build():
    print("=" * 70)
    print("MEMBANGUN PROPOSAL SKRIPSI FEB UKRIDA 2022 (BAB 1 - BAB 3)")
    print("=" * 70)

    # Read Source
    source_text = SOURCE_MD.read_text(encoding="utf-8")
    m1 = re.search(r"^# BAB 1", source_text, re.M)
    m4 = re.search(r"^# BAB 4", source_text, re.M)
    bab1_3_raw = source_text[m1.start():m4.start()]

    print(f"Sumber dimuat: {len(bab1_3_raw)} karakter, {bab1_3_raw.count(chr(10))} baris.")

    # ─────────────────────────────────────────────────────
    # BUILD MARKDOWN PROPOSAL
    # ─────────────────────────────────────────────────────
    print("\n[1/2] Menulis File Markdown: PROPOSAL_SKRIPSI_ARTHUR.md...")

    ref_keys_sorted = sorted(BIB_ENTRIES.keys(), key=lambda k: BIB_ENTRIES[k][1].get('author', ''))
    formatted_bib_md = []
    for k in ref_keys_sorted:
        formatted_bib_md.append(format_apa_reference(k, BIB_ENTRIES[k]))

    md_content = f"""# PROPOSAL SKRIPSI

<br><br>

# {JUDUL}

<br><br>

Diajukan kepada Program Studi Manajemen  
Untuk Menyusun Skripsi S1

<br><br><br>

**Diajukan Oleh:**  
**{PENULIS}**  
**({NIM})**  

<br><br><br><br>

**PROGRAM STUDI STRATA 1 MANAJEMEN**  
**FAKULTAS EKONOMI DAN BISNIS**  
**{UNIVERSITAS.upper()}**  
**{KOTA.upper()}**  
**{TAHUN}**  

---

<!-- Page Break -->

# PERSETUJUAN PROPOSAL SKRIPSI

<br>

| Data Mahasiswa | Keterangan |
|---|---|
| **Nama** | {PENULIS} ({JENIS_KELAMIN}) |
| **N.I.M.** | {NIM} |
| **Program Studi** | {PRODI} |
| **Konsentrasi** | {KONSENTRASI} |
| **Judul yang Diajukan** | {JUDUL} |

<br>

Proposal skripsi ini telah diperiksa dan disetujui untuk dipertahankan di hadapan Tim Penguji Seminar Proposal Skripsi {FAKULTAS} {UNIVERSITAS}.

<br><br>

{KOTA}, ________________ {TAHUN}  
**Menyetujui,**  

<br><br>

| Dosen Pembimbing | Dosen Pendamping |
|:---:|:---:|
| <br><br><br> **( {PEMBIMBING} )** | <br><br><br> **( ________________________ )** |

<br><br>

**Mengetahui,**  
**Ketua Program Studi Manajemen**  

<br><br><br>

**( {KAPRODI} )**  

---

<!-- Page Break -->

# DAFTAR ISI

- **HALAMAN JUDUL PROPOSAL** .................................................................... i
- **HALAMAN PERSETUJUAN PROPOSAL** .................................................... ii
- **DAFTAR ISI** ............................................................................................ iii
- **DAFTAR TABEL** ..................................................................................... iv
- **DAFTAR GAMBAR** ................................................................................... v
<br>
- **BAB 1 PENDAHULUAN** .......................................................................... 1
  - 1.1. Latar Belakang Penelitian .................................................................... 1
  - 1.2. Perumusan Masalah ............................................................................. 7
  - 1.3. Tujuan Penelitian ................................................................................ 8
  - 1.4. Manfaat Penelitian .............................................................................. 9
  - 1.5. Sistematika Penulisan Proposal Skripsi ................................................ 10
<br>
- **BAB 2 TINJAUAN PUSTAKA DAN PENGEMBANGAN HIPOTESIS** ................. 12
  - 2.1. Teori-Teori ........................................................................................... 12
    - 2.1.1. Teori Pemangku Kepentingan (*Stakeholder Theory*) ......................... 12
    - 2.1.2. Teori Legitimasi (*Legitimacy Theory*) ............................................. 14
    - 2.1.3. Teori Intermediasi Finansial (*Financial Intermediation Theory*) ......... 15
    - 2.1.4. Teori Sinyal (*Signaling Theory*) .................................................... 16
    - 2.1.5. Konsep Profitabilitas Bank dan *Return on Assets* (ROA) ................... 17
    - 2.1.6. Portofolio Kredit Hijau (*Green Financing*) ...................................... 18
    - 2.1.7. Risiko Kredit dan *Non-Performing Loan* (NPL) ............................... 20
    - 2.1.8. Kecukupan Modal dan *Capital Adequacy Ratio* (CAR) ..................... 22
  - 2.2. Penelitian Sebelumnya ....................................................................... 23
  - 2.3. Rerangka Penelitian ............................................................................ 30
  - 2.4. Pengembangan Hipotesis .................................................................... 32
    - 2.4.1. Pengaruh *Green Financing* terhadap *Return on Assets* (ROA) .......... 32
    - 2.4.2. Pengaruh *Non-Performing Loan* (NPL) terhadap ROA .................... 34
    - 2.4.3. Pengaruh *Capital Adequacy Ratio* (CAR) terhadap ROA ................... 36
    - 2.4.4. Pengaruh Simultan *Green Financing*, NPL, dan CAR terhadap ROA .... 37
<br>
- **BAB 3 METODE PENELITIAN** ................................................................. 39
  - 3.1. Jenis, Sumber, dan Teknik Pengumpulan Data ....................................... 39
  - 3.2. Populasi dan Sampel ............................................................................ 41
  - 3.3. Model Penelitian ................................................................................. 44
  - 3.4. Operasionalisasi Variabel .................................................................... 45
  - 3.5. Metode Analisis Data .......................................................................... 47
    - 3.5.1. Analisis Statistik Deskriptif ............................................................ 47
    - 3.5.2. Uji Stasioneritas Panel .................................................................. 48
    - 3.5.3. Penentuan Model Estimasi Regresi Panel .......................................... 49
    - 3.5.4. Uji Asumsi Klasik ......................................................................... 51
    - 3.5.5. Prosedur Estimasi EViews dan Robust Standard Errors ...................... 53
    - 3.5.6. Pengujian Hipotesis Statistik ......................................................... 54
<br>
- **DAFTAR PUSTAKA** ................................................................................. 57

---

<!-- Page Break -->

# DAFTAR TABEL

- **Tabel 2.1**  Matriks Sintesis Penelitian Terdahulu (Bagian 1: Penelitian 1–5) ............ 24
- **Tabel 2.2**  Matriks Sintesis Penelitian Terdahulu (Bagian 2: Penelitian 6–10) .......... 26
- **Tabel 2.3**  Matriks Sintesis Penelitian Terdahulu (Bagian 3: Penelitian 11–15) ......... 28
- **Tabel 3.1**  Daftar Entitas Sampel Penelitian Bank KBMI 4 ....................................... 43
- **Tabel 3.2**  Matriks Operasionalisasi Variabel Penelitian ......................................... 45

---

<!-- Page Break -->

# DAFTAR GAMBAR

- **Gambar 2.1** Rerangka Pemikiran Hubungan Kausalitas Variabel Penelitian ......... 31

---

<!-- Page Break -->

"""

    # Clean the body text for markdown
    cleaned_body = bab1_3_raw

    # Convert LaTeX tables to Markdown tables
    cleaned_body = convert_latex_tables_to_markdown(cleaned_body)

    # Clean headings
    cleaned_body = re.sub(r'#\s*BAB\s*1\s*\\+\[[^\]]*\]\s*PENDAHULUAN', '# BAB 1\n# PENDAHULUAN', cleaned_body)
    cleaned_body = re.sub(r'#\s*BAB\s*2\s*\\+\[[^\]]*\]\s*TINJAUAN PUSTAKA DAN PENGEMBANGAN HIPOTESIS', '# BAB 2\n# TINJAUAN PUSTAKA DAN PENGEMBANGAN HIPOTESIS', cleaned_body)
    cleaned_body = re.sub(r'#\s*BAB\s*3\s*\\+\[[^\]]*\]\s*METODE PENELITIAN', '# BAB 3\n# METODE PENELITIAN', cleaned_body)

    # Remove LaTeX boilerplate lines
    cleaned_body = re.sub(r'\\addcontentsline\{[^}]*\}\{[^}]*\}\{[^}]*\}', '', cleaned_body)
    cleaned_body = re.sub(r'\\addcontentsline\{[^}]*\}\{[^}]*\}', '', cleaned_body)
    cleaned_body = re.sub(r'\{BAB \d - [^}]*\}\}', '', cleaned_body)
    cleaned_body = re.sub(r'\\texorpdfstring\{[^}]*\}\{[^}]*\}', '', cleaned_body)
    cleaned_body = re.sub(r'\\setcounter\{[^}]*\}\{[^}]*\}', '', cleaned_body)
    cleaned_body = re.sub(r'\\vspace\{[^}]*\}', '', cleaned_body)
    cleaned_body = re.sub(r'\\begingroup', '', cleaned_body)
    cleaned_body = re.sub(r'\\endgroup', '', cleaned_body)
    cleaned_body = re.sub(r'\\singlespacing', '', cleaned_body)
    cleaned_body = re.sub(r'\\footnotesize', '', cleaned_body)
    cleaned_body = re.sub(r'\\small', '', cleaned_body)
    cleaned_body = re.sub(r'\\label\{[^}]*\}', '', cleaned_body)

    # Convert itemize and enumerate to clean markdown
    cleaned_body = re.sub(r'\\begin\{enumerate\}(?:\[[^\]]*\])?', '', cleaned_body)
    cleaned_body = re.sub(r'\\end\{enumerate\}', '', cleaned_body)
    cleaned_body = re.sub(r'\\begin\{itemize\}', '', cleaned_body)
    cleaned_body = re.sub(r'\\end\{itemize\}', '', cleaned_body)
    cleaned_body = re.sub(r'^\s*\\item\s*', '- ', cleaned_body, flags=re.MULTILINE)

    # Replace citations with readable format
    cleaned_body = clean_citations(cleaned_body)

    # Clean multiple blank lines
    cleaned_body = re.sub(r'\n{4,}', '\n\n', cleaned_body)

    md_content += cleaned_body
    md_content += "\n\n---\n\n<!-- Page Break -->\n\n# DAFTAR PUSTAKA\n\n"
    for ref in formatted_bib_md:
        md_content += f"{ref}\n\n"

    OUT_MD.write_text(md_content, encoding="utf-8")
    print(f"  Berhasil disimpan: {OUT_MD} ({len(md_content)//1024} KB)")

    # ─────────────────────────────────────────────────────
    # BUILD WORD (.DOCX) PROPOSAL
    # ─────────────────────────────────────────────────────
    print("\n[2/2] Menulis Dokumen Word: Proposal_Skripsi_Arthur.docx...")
    doc = Document()

    # Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)

    # ── SECTION 1: Bagian Awal (Nomor Romawi) ──
    sec1 = doc.sections[0]
    set_margins(sec1)
    add_footer_roman(sec1)

    print("  -> Halaman Judul Proposal...")
    build_judul_proposal(doc)
    doc.add_page_break()

    print("  -> Halaman Persetujuan Proposal...")
    build_persetujuan_proposal(doc)
    doc.add_page_break()

    print("  -> Daftar Isi Proposal...")
    build_daftar_isi_proposal(doc)
    doc.add_page_break()

    print("  -> Daftar Tabel Proposal...")
    build_daftar_tabel_proposal(doc)
    doc.add_page_break()

    print("  -> Daftar Gambar Proposal...")
    build_daftar_gambar_proposal(doc)

    # ── SECTION 2: Bagian Isi (Nomor Arab) ──
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(sec2)
    add_footer_arabic(sec2)

    print("  -> Menyusun BAB 1, BAB 2, BAB 3 ke dalam Word...")

    lines = bab1_3_raw.splitlines()
    i = 0
    in_enumerate = False
    enum_counter = 0

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        # Skip boilerplate
        if is_latex_noise(line) or not line:
            i += 1
            continue

        # ── Headings BAB ──
        if re.match(r'^#\s*BAB\s*\d', line):
            m = re.match(r'^#+\s*BAB\s*(\d+).*', line)
            bab_num = m.group(1) if m else "1"
            bab_titles = {
                "1": "PENDAHULUAN",
                "2": "TINJAUAN PUSTAKA DAN PENGEMBANGAN HIPOTESIS",
                "3": "METODE PENELITIAN",
            }
            title = bab_titles.get(bab_num, "")
            if bab_num != "1":
                doc.add_page_break()
            add_heading_bab(doc, bab_num, title)
            add_blank(doc)
            i += 1
            continue

        # ── Subheading H2 (## ) ──
        if line.startswith('## ') and not line.startswith('### '):
            title = clean_latex(line[3:])
            add_sub(doc, title, level=2)
            i += 1
            continue

        # ── Subheading H3 (### ) ──
        if line.startswith('### '):
            title = clean_latex(line[4:])
            add_sub(doc, title, level=3)
            i += 1
            continue

        # ── Subheading H4 (#### ) ──
        if line.startswith('#### '):
            title = clean_latex(line[5:])
            p = doc.add_paragraph()
            para_fmt(p, space_before=6, space_after=3)
            r = p.add_run(title)
            set_run_font(r, size=12, bold=True, italic=True)
            i += 1
            continue

        # ── Table block (\begin{table}) ──
        if re.match(r'\\begin\{table\}', line):
            j = i + 1
            table_lines = [line]
            while j < len(lines) and not re.match(r'\\end\{table\}', lines[j].strip()):
                table_lines.append(lines[j])
                j += 1
            if j < len(lines):
                table_lines.append(lines[j])
            table_raw = '\n'.join(table_lines)

            # Caption
            cap_m = re.search(r'\\caption\{([^}]+)\}', table_raw)
            caption = cap_m.group(1) if cap_m else "Tabel"

            # Source
            src_m = re.search(r'\\begin\{flushleft\}\s*\*?([^*]+)\*?\s*\\end\{flushleft\}', table_raw)
            source = src_m.group(1).strip() if src_m else ""

            # Rows
            tab_m = re.search(r'\\begin\{tabularx\}.*?\\end\{tabularx\}', table_raw, re.DOTALL)
            if tab_m:
                content = tab_m.group(0)
                row_lines = [rl.strip() for rl in content.split(r'\\') if '&' in rl]
                parsed_rows = []
                for rl in row_lines:
                    rl = re.sub(r'\\(hline|toprule|midrule|bottomrule)', '', rl).strip()
                    if not rl: continue
                    cells = [clean_latex(c) for c in rl.split('&')]
                    parsed_rows.append(cells)

                if parsed_rows:
                    headers = parsed_rows[0]
                    data_rows = parsed_rows[1:]
                    add_word_table_from_rows(doc, caption, headers, data_rows, source)

            i = j + 1
            continue

        # ── Enumerate / Itemize ──
        if re.match(r'\\begin\{enumerate', line):
            in_enumerate = True
            enum_counter = 0
            i += 1
            continue

        if re.match(r'\\end\{enumerate', line):
            in_enumerate = False
            enum_counter = 0
            i += 1
            continue

        if re.match(r'\\item', line):
            text_item = re.sub(r'^\s*\\item\s*(\[[^\]]*\])?\s*', '', line)
            j = i + 1
            while j < len(lines):
                nl = lines[j].strip()
                if (not nl or nl.startswith('\\item') or nl.startswith('\\begin') or
                        nl.startswith('\\end') or nl.startswith('#') or nl.startswith('%')):
                    break
                text_item += ' ' + nl
                j += 1
            text_item = clean_latex(text_item).strip()
            if text_item:
                if in_enumerate:
                    enum_counter += 1
                    add_numbered_item(doc, enum_counter, text_item)
                else:
                    add_bullet_item(doc, text_item)
            i = j
            continue

        # ── Markdown Numbered item (1. ) ──
        m_num = re.match(r'^(\d+)\.\s+(.+)', line)
        if m_num:
            num = m_num.group(1)
            text_item = clean_latex(m_num.group(2)).strip()
            j = i + 1
            while j < len(lines):
                nl = lines[j].strip()
                if not nl or re.match(r'^\d+\.', nl) or nl.startswith('#') or nl.startswith('\\'):
                    break
                text_item += ' ' + clean_latex(nl)
                j += 1
            if text_item:
                add_numbered_item(doc, num, text_item)
            i = j
            continue

        # ── Markdown Bullet item (- or •) ──
        if re.match(r'^[-*•]\s+', line):
            text_item = re.sub(r'^[-*•]\s+', '', line)
            text_item = clean_latex(text_item).strip()
            j = i + 1
            while j < len(lines):
                nl = lines[j].strip()
                if not nl or re.match(r'^[-*•]\s+', nl) or nl.startswith('#') or nl.startswith('\\'):
                    break
                text_item += ' ' + clean_latex(nl)
                j += 1
            if text_item:
                add_bullet_item(doc, text_item)
            i = j
            continue

        # ── Equations ──
        if re.match(r'\\begin\{equation', line):
            j = i + 1
            eq_lines = []
            while j < len(lines) and not re.match(r'\\end\{equation', lines[j].strip()):
                eq_lines.append(lines[j].strip())
                j += 1
            eq_text = clean_latex(' '.join(eq_lines))
            add_center(doc, eq_text, bold=True, size=11, space_before=6, space_after=6)
            i = j + 1
            continue

        # ── Paragraph Text ──
        para_lines = [line]
        j = i + 1
        while j < len(lines):
            nl = lines[j].strip()
            if (not nl or nl.startswith('#') or nl.startswith('%') or
                    re.match(r'^\s*\\(item|begin|end|vspace|vfill|hfill|newpage|clearpage)', nl) or
                    re.match(r'^---+$', nl) or re.match(r'^\d+\.\s+', nl) or
                    re.match(r'^[-*•]\s+', nl)):
                break
            if not is_latex_noise(nl):
                para_lines.append(nl)
            j += 1

        full_para = clean_latex(' '.join(para_lines))
        if full_para and len(full_para) > 3:
            add_body(doc, full_para, first_indent=True)

        i = j

    # ── DAFTAR PUSTAKA SECTION ──
    print("  -> Menambahkan DAFTAR PUSTAKA...")
    doc.add_page_break()
    add_heading_bab(doc, "", "DAFTAR PUSTAKA")
    add_blank(doc)

    for ref_str in formatted_bib_md:
        p = doc.add_paragraph()
        para_fmt(p, space_before=2, space_after=6, line_spacing=1.0,
                 first_indent=Cm(-1.25), left_indent=Cm(1.25))
        add_runs_with_inline(p, ref_str)

    # Save Document
    print(f"  -> Menyimpan file DOCX ke: {OUT_DOCX}...")
    doc.save(OUT_DOCX)
    size_kb = OUT_DOCX.stat().st_size // 1024
    print(f"\nSELESAI! Dokumen Word berhasil dibuat ({size_kb} KB).")
    print(f"File Output:\n1. {OUT_MD}\n2. {OUT_DOCX}\n")


if __name__ == "__main__":
    run_proposal_build()
