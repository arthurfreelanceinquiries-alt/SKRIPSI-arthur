import docx
import re
from pathlib import Path

docx_file = Path(r"z:\SKRIPSII\SKRIPSI-arthur\01_Naskah_Utama\Proposal_Arthur_PokemonTCG.docx")
doc = docx.Document(str(docx_file))

NEW_DOSBIM = "Dr. Fredella Colline, S.E., M.M., CFP®, PFM, CHCP-A"

def clean_and_replace(text):
    # Check if Diana is in text
    if "Diana Frederica" in text:
        # If it's the kata pengantar bullet
        if "selaku Dekan" in text:
            text = f"1. {NEW_DOSBIM}, selaku Dosen Pembimbing Skripsi, yang telah dengan luar biasa sabar, teliti, kritis, dan penuh dedikasi meluangkan waktu serta mencurahkan tenaga dan pikiran dalam membimbing, mengarahkan, dan menyempurnakan naskah proposal ini sejak tahap awal perumusan gagasan hingga penyusunan naskah komprehensif."
        elif "Thesis Advisor" in text:
            text = f"Thesis Advisor: {NEW_DOSBIM}"
        elif "Dosen Pembimbing:" in text:
            text = f"Dosen Pembimbing: {NEW_DOSBIM}"
        elif text.strip().startswith("(") and text.strip().endswith(")"):
            text = f"({NEW_DOSBIM})"
        else:
            text = NEW_DOSBIM
    if "0315088201" in text:
        text = text.replace("0315088201", "[NIDN_DOSEN]")
    return text

def process_paragraph(p):
    full_text = "".join(r.text for r in p.runs)
    if "Diana" in full_text or "0315088201" in full_text:
        new_text = clean_and_replace(full_text)
        if p.runs:
            p.runs[0].text = new_text
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.text = new_text

for para in doc.paragraphs:
    process_paragraph(para)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                process_paragraph(para)

doc.save(str(docx_file))
print("DOCX successfully patched!")
